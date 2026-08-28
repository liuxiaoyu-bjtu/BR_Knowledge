# -*- coding: utf-8 -*-
# U1-L4 P6 生成脚本：占位版 PPTX（照片类用占位框，框图类原生绘制，代码页浅色高亮，安全页警示红点缀）
# 视觉规范 §8.12：知识蓝 / 实践青绿 / 思考评价橘黄 / 安全警示红
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
try:
    from pptx.enum.dml import MSO_DASH_STYLE
except Exception:
    MSO_DASH_STYLE = None
from pptx.oxml.ns import qn

# ---------- 基础常量 ----------
PAGE_W = Inches(13.333)
PAGE_H = Inches(7.5)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
BLUE    = RGBColor(0x25, 0x63, 0xEB)   # 科技蓝 主色
DARK    = RGBColor(0x1F, 0x29, 0x37)   # 主文字
GRAY    = RGBColor(0x60, 0x6A, 0x78)   # 副文字
LIGHT   = RGBColor(0xF4, 0xF6, 0xF9)   # 浅底/代码底
TEAL    = RGBColor(0x0E, 0xA5, 0xA4)   # 实践操作主色
TEALBG  = RGBColor(0xCC, 0xFB, 0xF1)   # 实践代码高亮
AMBER   = RGBColor(0xFF, 0xF7, 0xE6)   # 思考/评价浅底
AMBERBAR= RGBColor(0xF5, 0x9E, 0x0B)   # 思考/评价主色
PH_FILL = RGBColor(0xEE, 0xF1, 0xF5)   # 占位框底
PH_LINE = RGBColor(0xB8, 0xC0, 0xCC)   # 占位框边
SOFT    = RGBColor(0xE7, 0xEC, 0xF5)   # 浅蓝卡片
CARD    = RGBColor(0xF0, 0xF4, 0xFC)   # 浅蓝卡
WARN    = RGBColor(0xE5, 0x48, 0x4D)   # 警示红（仅安全页点缀）
WARNBG  = RGBColor(0xFD, 0xE7, 0xE7)   # 警示红浅底
FONT    = "Microsoft YaHei"
MONO    = "Consolas"
TOTAL   = 26

OUT = r"d:/A_Project/BoosterRobotics/知识库/05-internal-materials/vocational-planning/03_X课程项目/PPT/U1-L4_头部步态与预定义动作_Codex方案二审查版.pptx"

prs = Presentation()
prs.slide_width = PAGE_W
prs.slide_height = PAGE_H
BLANK = prs.slide_layouts[6]

# ---------- 字体助手 ----------
def ea(run, font=FONT):
    rPr = run._r.get_or_add_rPr()
    rPr.set(qn('a:ea'), font)
    rPr.set(qn('a:cs'), font)

def style_run(run, size=14, color=DARK, bold=False, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    ea(run, font)

def set_text(shape, text, size=14, color=DARK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, font=FONT, line_spacing=1.12, runs=None):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.06)
    tf.margin_bottom = Inches(0.06)
    if runs is None:
        lines = text.split("\n")
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line_spacing
            r = p.add_run()
            r.text = line
            style_run(r, size, color, bold, font)
    else:
        for i, (txt, sz, col, b, fnt) in enumerate(runs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line_spacing
            r = p.add_run()
            r.text = txt
            style_run(r, sz, col, b, fnt)

# ---------- 形状助手 ----------
def shape(slide, kind, l, t, w, h, fill=None, line=None, line_w=1.0, dash=False):
    sp = slide.shapes.add_shape(kind, l, t, w, h)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
        if dash and MSO_DASH_STYLE is not None:
            try:
                sp.line.dash_style = MSO_DASH_STYLE.DASH
            except Exception:
                pass
    sp.shadow.inherit = False
    return sp

def card(slide, l, t, w, h, text, fill=SOFT, line=None, line_w=1.0,
         size=14, color=DARK, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
         font=FONT, dashed=False):
    sp = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h, fill=fill, line=line or fill, line_w=line_w, dash=dashed)
    set_text(sp, text, size=size, color=color, bold=bold, align=align, anchor=anchor, font=font)
    return sp

def arrow(slide, kind, l, t, w, h, fill=BLUE):
    sp = slide.shapes.add_shape(kind, l, t, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp

# ---------- 页面通用元素 ----------
def new_slide():
    return prs.slides.add_slide(BLANK)

def bg(slide):
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, PAGE_W, PAGE_H, fill=WHITE)

def section_tag(slide, text, color=BLUE):
    if not text:
        return
    w = Inches(min(7.0, 0.35 + 0.18 * len(text)))
    sp = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(0.42), w, Inches(0.48),
               fill=color, line=color)
    set_text(sp, text, size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def title_block(slide, title, subtitle=None):
    set_text(shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.02), Inches(12.0), Inches(0.62), fill=None, line=None),
             title, size=32, color=DARK, bold=True, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        set_text(shape(slide, MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.66), Inches(12.0), Inches(0.45), fill=None, line=None),
             subtitle if subtitle.lstrip().startswith("——") else "—— " + subtitle, size=18, color=GRAY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)

def page_no(slide, n):
    set_text(shape(slide, MSO_SHAPE.RECTANGLE, Inches(11.9), Inches(7.0), Inches(1.2), Inches(0.34), fill=None, line=None),
             "%d / %d" % (n, TOTAL), size=11, color=GRAY, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

def placeholder(slide, l, t, w, h, label=""):
    sp = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h, fill=PH_FILL, line=PH_LINE, line_w=1.5, dash=True)
    txt = "【配图占位】"
    if label:
        txt += "\n" + label
    set_text(sp, txt, size=14, color=GRAY, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def bullets(slide, l, t, w, h, items, size=16, color=DARK, gap=1.18, runs_list=None):
    sp = shape(slide, MSO_SHAPE.RECTANGLE, l, t, w, h, fill=None, line=None)
    tf = sp.text_frame
    tf.word_wrap = True
    if runs_list is None:
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = gap
            p.space_after = Pt(6)
            r = p.add_run()
            r.text = "• " + it
            style_run(r, size, color, False)
    else:
        for i, (txt, sz, col, b) in enumerate(runs_list):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = gap
            p.space_after = Pt(6)
            r = p.add_run()
            r.text = txt
            style_run(r, sz, col, b)

def code_block(slide, l, t, w, lines, highlight=None, size=13):
    highlight = highlight or set()
    line_h = Inches(0.36)
    n = len(lines)
    h = Emu(int(line_h * (n + 1)) + int(Inches(0.24)))
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h, fill=LIGHT, line=PH_LINE, line_w=1.0)
    shape(slide, MSO_SHAPE.RECTANGLE, l, t, Inches(0.09), h, fill=TEAL, line=None)
    for i in highlight:
        y = Emu(int(t) + int(Inches(0.12)) + i * int(line_h))
        shape(slide, MSO_SHAPE.RECTANGLE, l, y, w, line_h, fill=TEALBG, line=None)
    tb = shape(slide, MSO_SHAPE.RECTANGLE, l, t, w, h, fill=None, line=None)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.0
        r = p.add_run()
        r.text = ln if ln else " "
        style_run(r, size, DARK, False, MONO)
    return h

def hint_card(slide, l, t, w, h, kind, text):
    accent = TEAL if kind == "动手试试" else (AMBERBAR if kind == "想一想" else BLUE)
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h, fill=CARD, line=accent, line_w=1.0)
    shape(slide, MSO_SHAPE.RECTANGLE, l, t, Inches(0.10), h, fill=accent, line=None)
    set_text(shape(slide, MSO_SHAPE.RECTANGLE, l + Inches(0.25), t + Inches(0.06), w - Inches(0.4), Inches(0.4), fill=None, line=None),
             "【%s】" % kind, size=14, color=accent, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    set_text(shape(slide, MSO_SHAPE.RECTANGLE, l + Inches(0.25), t + Inches(0.46), w - Inches(0.45), h - Inches(0.52), fill=None, line=None),
             text, size=13, color=DARK, anchor=MSO_ANCHOR.TOP)

def warn_bar(slide, l, t, w, text):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, Inches(0.5), fill=WARN, line=WARN)
    set_text(shape(slide, MSO_SHAPE.RECTANGLE, l, t, w, Inches(0.5), fill=None, line=None),
             text, size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def concept_tag(slide, text):
    set_text(shape(slide, MSO_SHAPE.RECTANGLE, Inches(9.3), Inches(6.55), Inches(3.4), Inches(0.4), fill=None, line=None),
             text, size=13, color=BLUE, bold=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

# ---------- L4 专属原生表格 ----------
def draw_head_table(s):
    rows, cols = 3, 4
    gf = s.shapes.add_table(rows, cols, Inches(0.9), Inches(2.7), Inches(11.5), Inches(3.0))
    tbl = gf.table
    tbl.columns[0].width = Inches(2.4); tbl.columns[1].width = Inches(3.0)
    tbl.columns[2].width = Inches(3.4); tbl.columns[3].width = Inches(2.7)
    data = [
        ("角度", "名称", "正值方向", "物理限位（约）"),
        ("pitch", "俯仰角", "低头（看向地面）", "-0.3 ~ 1.0 rad"),
        ("yaw", "偏航角", "头向左转", "-0.785 ~ 0.785 rad"),
    ]
    for r in range(rows):
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.text = data[r][c]
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE if r == 0 else (LIGHT if c == 0 else WHITE)
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            run = para.runs[0]
            style_run(run, 14, WHITE if r == 0 else DARK, bold=(r == 0 or c == 0), font=(MONO if c == 0 else FONT))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)

def draw_head_params_table(s):
    rows, cols = 3, 4
    gf = s.shapes.add_table(rows, cols, Inches(0.9), Inches(2.7), Inches(11.5), Inches(3.0))
    tbl = gf.table
    tbl.columns[0].width = Inches(2.4); tbl.columns[1].width = Inches(3.0)
    tbl.columns[2].width = Inches(3.4); tbl.columns[3].width = Inches(2.7)
    data = [
        ("参数", "含义", "正方向", "模式约束"),
        ("pitch", "头部俯仰", "正值 = 低头", "仅在 walk 模式下有效"),
        ("yaw", "头部偏航", "正值 = 头向左转", "仅在 walk 模式下有效"),
    ]
    for r in range(rows):
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.text = data[r][c]
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE if r == 0 else (LIGHT if c == 0 else WHITE)
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            run = para.runs[0]
            style_run(run, 14, WHITE if r == 0 else DARK, bold=(r == 0 or c == 0), font=(MONO if c == 0 else FONT))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)

def draw_head_action_table(s):
    rows, cols = 4, 3
    gf = s.shapes.add_table(rows, cols, Inches(0.9), Inches(2.7), Inches(11.5), Inches(3.0))
    tbl = gf.table
    tbl.columns[0].width = Inches(5.4); tbl.columns[1].width = Inches(3.4); tbl.columns[2].width = Inches(2.7)
    data = [
        ("指令", "动作", "方向"),
        ("set_head_angle(0.3, 0)", "低头", "看向地面"),
        ("set_head_angle(0, 0.5)", "头向左转", "向左 0.5 rad"),
        ("set_head_angle(0, 0)", "回正", "中位"),
    ]
    for r in range(rows):
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.text = data[r][c]
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE if r == 0 else (LIGHT if c == 0 else WHITE)
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            run = para.runs[0]
            style_run(run, 13, WHITE if r == 0 else DARK, bold=(r == 0 or c == 0), font=(MONO if c == 0 else FONT))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)

def draw_gait_table(s):
    rows, cols = 3, 3
    gf = s.shapes.add_table(rows, cols, Inches(0.9), Inches(2.8), Inches(11.5), Inches(2.8))
    tbl = gf.table
    tbl.columns[0].width = Inches(3.4); tbl.columns[1].width = Inches(4.05); tbl.columns[2].width = Inches(4.05)
    data = [
        ("维度", "soccer（足球步态）", "default（标准步态）"),
        ("适用", "踢球等前倾动作", "日常平稳行走"),
        ("重心 / 特点", "更靠前、利于踢球", "居中、更稳"),
    ]
    for r in range(rows):
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.text = data[r][c]
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE if r == 0 else (LIGHT if c == 0 else WHITE)
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.CENTER
            run = para.runs[0]
            style_run(run, 14, WHITE if r == 0 else DARK, bold=(r == 0 or c == 0))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)

def draw_enter_walk_timing_l4(s):
    y = Inches(3.2); bw = Inches(1.9); bh = Inches(1.7)
    xs = [Inches(0.7), Inches(2.7), Inches(4.7), Inches(6.7), Inches(8.7), Inches(10.7)]
    labels = ["① 进入 prepare\n并手动拉起站稳", "② 人工确认\n仅首次·直立·空旷·教师监督",
              "③ 进入 walk\n发头部/步态/动作", "④ 回正头部\nset_head_angle 0,0",
              "⑤ 回 prepare\n站立稳定", "⑥ 手动趴姿\n→ damping（人工）"]
    for i, (lb, x) in enumerate(zip(labels, xs)):
        fillc = WARN if i == 1 else SOFT
        card(s, x, y, bw, bh, lb, fill=fillc, size=12,
             color=(WHITE if i == 1 else DARK), bold=(i == 1))
    for i in range(5):
        arrow(s, MSO_SHAPE.RIGHT_ARROW, Emu(int(xs[i]) + int(bw)), y + Inches(0.6), Inches(0.3), Inches(0.4))

# ---------- 页面生成 ----------
def page_accent(n):
    if n in {25, 26}:
        return AMBERBAR
    if n in {16, 17, 18, 19, 20, 21, 22, 23, 24}:
        return TEAL
    return BLUE

def make(slide, n, section, title, subtitle, build):
    bg(slide)
    if section:
        section_tag(slide, section, page_accent(n))
    else:
        section_tag(slide, "单元 U1 · 人形机器人初探与运动控制")
    title_block(slide, title, subtitle)
    build(slide)
    page_no(slide, n)

# 1 封面
s = new_slide(); bg(s)
section_tag(s, "第四堂人形机器人课")
set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.3), fill=None, line=None),
         "头部、步态与预定义动作", size=44, color=DARK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(3.7), Inches(11.5), Inches(0.6), fill=None, line=None),
         "—— 让 K1 会看、会走、会做", size=22, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)
shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(4.5), Inches(2.2), Inches(0.10), fill=BLUE, line=None)
placeholder(s, Inches(0.9), Inches(4.9), Inches(11.5), Inches(1.7), "课程主视觉 / K1 转动头部 + 行走场景 或 头部坐标系示意图")
set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(6.8), Inches(11.5), Inches(0.4), fill=None, line=None),
         "课时 L4 · 主讲人：__________", size=14, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)
page_no(s, 1)

# 2 本课目标
def p2(s):
    bullets(s, Inches(0.9), Inches(2.5), Inches(11.5), Inches(3.5), [
        "认识头部坐标系 pitch / yaw，会用 set_head_angle 控制头部姿态",
        "掌握 set_gait 切换足球 / 标准步态，对比行走差异",
        "用 do_action 调用预定义动作，并用 TaskHandle 确认任务终态，体验开环组合行为",
    ], size=20)
make(new_slide(), 2, "导学", "本课目标", "—— 一节课我们要带走什么", p2)

# 3 互动导入
def p3(s):
    placeholder(s, Inches(7.4), Inches(2.4), Inches(5.2), Inches(3.4), "K1 站立待发 / 头部转动 场景图")
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.4), Inches(6.2), Inches(1.4), fill=None, line=None),
             "回顾：上节课我们让 K1 按指令走起来了。", size=18, color=DARK, bold=True)
    hint_card(s, Inches(0.9), Inches(4.0), Inches(6.2), Inches(1.8), "想一想",
              "想让它走路时脑袋转向一边、\n看看那边的东西，程序该怎么写？\n先猜一下，下面揭晓。")
make(new_slide(), 3, "第一部分 · 头部坐标系与角度指令", "从“能走”到“会看、会做”", "—— 脑袋和走路姿势也能编程吗？", p3)

# 4 什么是头部坐标系
def p4(s):
    draw_head_table(s)
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(6.1), Inches(11.5), Inches(0.8), fill=None, line=None),
             "头部自身坐标系：pitch 管低头抬头，yaw 管头左右转；单位都是弧度（rad）。", size=14, color=GRAY)
make(new_slide(), 4, "第一部分 · 头部坐标系与角度指令", "什么是头部坐标系", "—— 以机器人头部为参照的两个角度", p4)

# 5 两个角度的正方向
def p5(s):
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.4), Inches(6.2), Inches(3.0), fill=None, line=None),
             "两个角度的正方向约定：\n\n"
             "· pitch 正值 = 低头（看向地面）\n"
             "· yaw 正值 = 头向左转\n\n"
             "注意：yaw 与 L3 机身 yaw 同名，但一个是整机转向、一个是脑袋转头。", size=16, color=DARK)
    placeholder(s, Inches(7.4), Inches(2.4), Inches(5.2), Inches(3.4), "头部正方向标注图（pitch 低头 / yaw 左转）")
make(new_slide(), 5, "第一部分 · 头部坐标系与角度指令", "两个角度的正方向", "—— pitch 正 = 低头 · yaw 正 = 头向左转", p5)

# 6 set_head_angle 参数与约束
def p6(s):
    draw_head_params_table(s)
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(6.1), Inches(11.5), Inches(0.8), fill=None, line=None),
             "set_head_angle(pitch, yaw) 同时设置、立即生效；但只在 walk 模式下有效，否则报错。", size=14, color=GRAY)
make(new_slide(), 6, "第一部分 · 头部坐标系与角度指令", "set_head_angle 的参数与约束", "—— 一个接口，两个角度，只听 walk 的", p6)

# 7 头部视野与运动
def p7(s):
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.4), Inches(6.2), Inches(3.0), fill=None, line=None),
             "头部相机是 K1“看世界”的主要来源。\n\n"
             "行走中改变 yaw 让头部左右转动，就能比固定朝前看到更宽的视野——\n"
             "这也是后续“看到目标再行动”的第一步。", size=16, color=DARK)
    placeholder(s, Inches(7.4), Inches(2.4), Inches(5.2), Inches(3.4), "头部视野扫描示意图（固定朝前 vs 左右转动）")
make(new_slide(), 7, "第一部分 · 头部坐标系与角度指令", "头部视野与运动", "—— 边走边转头，能看更宽", p7)

# 8 三个头部动作对照
def p8(s):
    draw_head_action_table(s)
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(6.1), Inches(11.5), Inches(0.8), fill=None, line=None),
             "低头 / 转头 / 回正三个动作怎么写，一表记牢；均在 walk 模式下发送。", size=14, color=GRAY)
make(new_slide(), 8, "第一部分 · 头部坐标系与角度指令", "三个头部动作对照", "—— 低头 / 转头 / 回正怎么写", p8)

# 9 互动：引出步态与动作库
def p9(s):
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.8), Inches(11.3), Inches(1.6), fill=None, line=None),
             "除了脑袋，还能怎么“走”和“做”？\n—— 它的走路姿势可调吗？厂家还准备好了哪些现成动作？", size=30, color=DARK, bold=True)
    placeholder(s, Inches(9.0), Inches(1.2), Inches(3.3), Inches(2.0), "K1 不同步态 / 做动作 示意")
make(new_slide(), 9, "第二部分 · 步态与预定义动作", "除了脑袋，还能怎么“走”和“做”", "—— 引出步态与动作库", p9)

# 10 步态是什么
def p10(s):
    draw_gait_table(s)
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(6.0), Inches(11.5), Inches(0.9), fill=None, line=None),
             "set_gait(\"soccer\" | \"default\")：步态决定“怎么迈步”，本课对比两种步态的行走差异。", size=14, color=GRAY)
make(new_slide(), 10, "第二部分 · 步态与预定义动作", "步态是什么", "—— soccer 踢球步 vs default 标准步", p10)

# 11 预定义动作库
def p11(s):
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.4), Inches(6.2), Inches(3.2), fill=None, line=None),
             "预定义动作库 = 厂家把常用肢体行为固化好，可通过 do_action(id) 一行调用。\n\n"
             "调用会立即返回一个 TaskHandle（任务句柄），用来跟踪这次动作任务是否结束；\n"
             "不必关心底层关节轨迹，只关心「做什么、何时做完」。", size=16, color=DARK)
    placeholder(s, Inches(7.4), Inches(2.4), Inches(5.2), Inches(3.0), "预定义动作缩略图（以 list_actions 读到的动作为准）")
    concept_tag(s, "核心概念：动作库 + TaskHandle")
make(new_slide(), 11, "第二部分 · 步态与预定义动作", "预定义动作库", "—— do_action(id)，一行调用现成行为", p11)

# 12 看清武器库（概念页，代码见实验四 4.1）
def p12(s):
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.4), Inches(6.2), Inches(3.2), fill=None, line=None),
             "调用动作前，先看清“武器库”：\n\n"
             "· `list_actions()` 列出所有预定义动作及其 id、type、duration、interruptible\n"
             "· `do_action(id)` 用 id 指定动作（id 比名字稳定），返回 `TaskHandle`\n"
             "· `list_gaits()` 查看支持的步态\n\n"
             "只读取、不移动；具体代码在实验四（4.1）调用前查看。", size=16, color=DARK)
    placeholder(s, Inches(7.4), Inches(2.4), Inches(5.2), Inches(3.0), "动作库清单截图 / 动作缩略图")
make(new_slide(), 12, "第二部分 · 步态与预定义动作", "看清手里的“武器库”", "—— 先列清单，再调用", p12)

# 13 本课认知主线
def p13(s):
    card(s, Inches(0.9), Inches(2.8), Inches(3.6), Inches(3.0),
         "① 头部坐标系\n\npitch 低头\nyaw 头左转\n\n用 set_head_angle", fill=CARD, size=16, color=DARK)
    card(s, Inches(4.85), Inches(2.8), Inches(3.6), Inches(3.0),
         "② 步态切换\n\nsoccer / default\n\n用 set_gait", fill=CARD, size=16, color=DARK)
    card(s, Inches(8.8), Inches(2.8), Inches(3.6), Inches(3.0),
         "③ 预定义动作\n\ndo_action(id)\n\n动作库一行调用", fill=CARD, size=16, color=DARK)
make(new_slide(), 13, "第二部分 · 步态与预定义动作", "本课认知主线", "—— 头部 + 步态 + 动作 = 组合行为", p13)

# 14 过渡：安全
def p14(s):
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.8), Inches(11.3), Inches(1.6), fill=None, line=None),
             "动手前，先守规矩\n—— 接下来 K1 会动头部、换步态、做大幅动作，比“原地切换”更需谨慎", size=30, color=DARK, bold=True)
    placeholder(s, Inches(9.0), Inches(1.2), Inches(3.3), Inches(2.0), "空旷场地 + 监督 场景剪影")
make(new_slide(), 14, "第三部分 · 头部·步态·动作实验与安全", "动手前，先守规矩", "—— 真机移动与大幅动作的安全底线", p14)

# 15 实验安全规范（警示红点缀）
def p15(s):
    warn_bar(s, Inches(0.9), Inches(2.4), Inches(11.5), "⚠ 安全提示：真机移动与大幅动作，安全优先")
    bullets(s, Inches(0.9), Inches(3.3), Inches(11.5), Inches(3.0), [
        "① 场地空旷、无易撞物，远离台阶与电源",
        "② 教师与同伴在场，保持安全距离",
        "③ 明确急停与收尾手势，人人知晓",
        "④ 大幅动作前确认周围空间充足",
    ], size=18, gap=1.3, color=DARK)
make(new_slide(), 15, "第三部分 · 头部·步态·动作实验与安全", "实验安全规范", "—— 空旷 · 在场 · 有预案", p15)

# 16 代码：连接 K1
def p16(s):
    code_block(s, Inches(0.9), Inches(2.6), Inches(11.5),
               ["# 对应 notebook 2.1（cell-7）",
                "from boosteros.robots.booster import BoosterRobot",
                "robot = BoosterRobot()      # 建立连接（仅连接，不移动）"],
               highlight={1, 2})
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(5.2), Inches(11.5), Inches(1.2), fill=None, line=None),
             "先连上机器人；动作库与步态清单在实验四（4.1）调用前再用 `list_actions()` 查看，不必提前统一确认模式。", size=15, color=GRAY)
make(new_slide(), 16, "第三部分 · 头部·步态·动作实验与安全", "实验准备：连接 K1", "—— 先连上，进实验再看清清单", p16)

# 17 进入行走态时序（仅首次人工窗口，红点缀）
def p17(s):
    warn_bar(s, Inches(0.9), Inches(2.4), Inches(11.5), "⚠ 人工确认仅首次：直立站稳 · 周围空旷 · 教师监督")
    draw_enter_walk_timing_l4(s)
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(5.6), Inches(11.5), Inches(0.9), fill=None, line=None),
             "六步时序，② 人工确认仅首次不可跳过；其余步骤只在代码前注明所需模式，不重复窗口。对应 notebook 2.3 首次进入 walk（cell-9 / cell-10）。",
             size=14, color=GRAY)
make(new_slide(), 17, "第三部分 · 头部·步态·动作实验与安全", "进入行走态（仅首次需人工窗口）", "—— 确认后才发 walk，后续只注模式", p17)

# 18 代码：头部控制
def p18(s):
    code_block(s, Inches(0.9), Inches(2.6), Inches(11.5),
               ["# 对应 notebook 2.3–2.5（cell-10 / cell-12 / cell-14），机器人处于 walk 模式",
                "robot.set_head_angle(0.3, 0)    # 低头约 0.3 rad",
                "robot.set_head_angle(0, 0.5)    # 头向左转 0.5 rad",
                "robot.set_head_angle(0, 0)      # 回正"],
               highlight={1, 2, 3})
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(5.6), Inches(11.5), Inches(0.9), fill=None, line=None),
             "机器人保持在 walk 模式；先不发速度，仅调头部，观察视野变化。运行前确认已人工窗口进入 walk。", size=15, color=GRAY)
make(new_slide(), 18, "第三部分 · 头部·步态·动作实验与安全", "代码 · 头部控制", "—— set_head_angle(0.3, 0) 等", p18)

# 19 代码：步态切换
def p19(s):
    code_block(s, Inches(0.9), Inches(2.6), Inches(11.5),
               ["# 对应 notebook 3.1–3.3（cell-17 / cell-19 / cell-21），机器人处于 walk 模式",
                "robot.set_gait(\"soccer\")           # 切换足球步态",
                "robot.set_velocity(0.3, 0, 0)      # 前进 0.3 m/s",
                "# 约 2 秒后",
                "robot.set_velocity(0, 0, 0)         # 停止",
                "robot.set_gait(\"default\")          # 切回标准步态，同样前进对比"],
               highlight={1, 2, 6})
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(6.0), Inches(11.5), Inches(0.6), fill=None, line=None),
             "本环节在 walk 模式下切换 gait 并对比（实验二头部控制即在该模式下）；保持低速短时，停止后回 prepare（见 3.4）。", size=14, color=GRAY)
make(new_slide(), 19, "第三部分 · 头部·步态·动作实验与安全", "代码 · 步态切换", "—— set_gait(\"soccer\") / \"default\"", p19)

# 20 代码：预定义动作（含 4.1 看清单）
def p20(s):
    code_block(s, Inches(0.9), Inches(2.3), Inches(11.5),
               ["actions = robot.list_actions()       # 先看清单（只读）",
                "print([(a.id, a.type) for a in actions])",
                "print(\"支持的步态：\", robot.list_gaits())",
                "act = actions[0]; handle = robot.do_action(act.id)  # walk 模式",
                "print(\"status =\", handle.status); handle.wait()",
                "print(\"done =\", handle.done(), \"final =\", handle.status)",
                "upper = next(a for a in actions if a.type == \"upper_body\")",
                "whole = next(a for a in actions if a.type == \"whole_body\")",
                "h1 = robot.do_action(upper.id); h1.wait()",
                "h2 = robot.do_action(whole.id); h2.wait()"],
               highlight={0, 3, 4, 8, 9})
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(6.7), Inches(11.5), Inches(0.5), fill=None, line=None),
             "do_action 必须在 walk 模式下调用；幅度大，确认周围空旷、教师监督后再运行。运行后先回 prepare（4.4），再手动趴姿→damping。", size=13, color=GRAY)
make(new_slide(), 20, "第三部分 · 头部·步态·动作实验与安全", "代码 · 预定义动作", "—— 看清清单后 do_action(id)，并用 TaskHandle 确认终态", p20)

# 21 代码：停止与安全收尾（警示红）
def p21(s):
    warn_bar(s, Inches(0.9), Inches(2.4), Inches(11.5), "⚠ 头部回正 → walk→prepare；动作实验结束先回 prepare（4.4），再手动趴姿→damping")
    code_block(s, Inches(0.9), Inches(3.15), Inches(11.5),
               ["# 对应 notebook 2.5 / 3.4 / 4.4（cell-14 / cell-23 / cell-32）",
                "robot.set_head_angle(0, 0)      # 头部回正",
                "robot.set_mode(\"prepare\")       # 安全停止，回到站立稳定",
                "# 确认已站立稳定后，按约定手动放置趴姿 → damping → 关机（严禁程序直进 damping）"],
               highlight={1, 2})
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(5.9), Inches(11.5), Inches(0.8), fill=None, line=None),
             "头部实验先回正再 walk→prepare；步态实验先归零再 prepare；动作实验结束先回 prepare（4.4），再手动趴姿→damping，严禁程序直接 walk→damping。", size=14, color=WARN, bold=True)
make(new_slide(), 21, "第三部分 · 头部·步态·动作实验与安全", "代码 · 停止与安全收尾", "—— 先回正，再 walk→prepare，最后 damping", p21)

# 22 观察记录任务
def p22(s):
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.4), fill=None, line=None),
             "不写代码：每组按指令操作头部 / 步态 / 动作，记录实际现象与偏差，填入实训手册 §六 观察表"
             "（开环：发指令、测现象、做记录）。", size=17, color=DARK, bold=True)
    hint_card(s, Inches(0.9), Inches(4.0), Inches(11.5), Inches(2.4), "动手试试",
              "观察记录要点：\n\n"
              "· 头部：低头 / 转头后相机视野怎么变？\n"
              "· 步态：同样 0.3 m/s 走 2 秒，soccer 与 default 姿态 / 稳定感有何不同？\n"
              "· 动作：每个 do_action 做了什么、幅度多大？把现象写成具体描述，同桌交叉核对。")
make(new_slide(), 22, "第三部分 · 头部·步态·动作实验与安全", "观察记录任务", "—— 把“预期”和“实际”的差记下来", p22)

# 23 拓展：边走边转头扫描
def p23(s):
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.4), Inches(6.2), Inches(3.2), fill=None, line=None),
             "把头部控制与 L3 速度控制组合起来，可形成“前进同时头部左右扫视”的组合行为：\n\n"
             "· 先写指令顺序（先做什么、再做什么、何时回正）\n"
             "· 标注每步所需模式（哪些必须在 walk 下）\n"
             "· 列出至少 2 条安全注意点", size=15, color=DARK)
    placeholder(s, Inches(7.4), Inches(2.4), Inches(5.2), Inches(3.2), "组合行为示意图（前进 + 头部左右扫视）")
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(7.4), Inches(5.8), Inches(5.2), Inches(0.5), fill=None, line=None),
             "详细图示见附件《K1 综合行为设计任务卡》", size=13, color=BLUE, bold=True)
make(new_slide(), 23, "第四部分 · 拓展与评价", "拓展：边走边转头扫描", "—— 见附件《K1 综合行为设计任务卡》", p23)

# 24 分组实操任务卡
def p24(s):
    hint_card(s, Inches(0.9), Inches(2.5), Inches(11.5), Inches(3.8), "动手试试",
              "分组轮换，在监督下完成一次组合行为：\n\n"
              "· 每组：prepare → 确认 → 进 walk → 低头/转头/回正 → 切换步态 → 调动作 → 停止收尾\n"
              "· 每组指定 1 名安全员，全程盯守\n"
              "· 教师先完整示范一遍，再放手让学生操作\n"
              "· 每个动作运行前确认周围空旷、站位安全；动作幅度大须留足空间")
make(new_slide(), 24, "第四部分 · 拓展与评价", "分组实操任务卡", "—— 在监督下完成一次组合行为", p24)

# 25 本课小结（三卡）
def p25(s):
    card(s, Inches(0.9), Inches(2.8), Inches(3.6), Inches(3.0),
         "① 头部坐标系\n\npitch 低头\nyaw 头左转\n\n与 L3 机身 yaw 不同", fill=CARD, size=16, color=DARK)
    card(s, Inches(4.85), Inches(2.8), Inches(3.6), Inches(3.0),
         "② 步态切换\n\nsoccer / default\n\n用 set_gait", fill=CARD, size=16, color=DARK)
    card(s, Inches(8.8), Inches(2.8), Inches(3.6), Inches(3.0),
         "③ 预定义动作库\n\ndo_action(id)\n\n一行调用现成行为", fill=CARD, size=16, color=DARK)
make(new_slide(), 25, "第四部分 · 拓展与评价", "本课小结", "—— 带走三件东西", p25)

# 26 学习评价
def p26(s):
    set_text(shape(s, MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.3), Inches(11.5), Inches(0.5), fill=None, line=None),
             "带走三件东西：① 头部坐标系　② 步态切换　③ 预定义动作库", size=16, color=DARK, bold=True)
    card(s, Inches(0.9), Inches(3.1), Inches(11.5), Inches(0.6), "学习评价（与实训手册一致）", fill=AMBERBAR, line=AMBERBAR, color=WHITE, bold=True, size=15)
    bullets(s, Inches(1.0), Inches(3.95), Inches(11.3), Inches(3.0), [
        "判断：set_head_angle(pitch, yaw) 只能在 walk 模式下调用，在 damping / prepare 下调用会报错。（　）",
        "填空：切换步态用 set_gait(____)；若当前不在 walk 模式，该设置会 ______（填“立即生效”或“记录并在下次进入 walk 时生效”）。",
        "填空：头部俯仰角 pitch 的正值表示 ______（填“低头”或“抬头”）。",
        "选择：用于在 walk 模式下控制头部姿态的接口是（　）。A. set_velocity　B. set_gait　C. set_head_angle　D. do_action",
        "选择：do_action 指定动作时应传入（　）。A. 动作名字字符串　B. list_actions() 返回的 id　C. 动作 duration　D. 任意整数",
    ], size=13, gap=1.02)
make(new_slide(), 26, "第四部分 · 拓展与评价", "学习评价", "—— 测一测，带得走多少", p26)

prs.save(OUT)
print("DONE: %s  slides=%d" % (OUT, len(prs.slides._sldIdLst)))
